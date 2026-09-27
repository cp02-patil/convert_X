import io
from pptx import Presentation
import pptx as tc

def pptx_to_text(data):
    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"[Slide {i}]")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts)

def pptx_to_pdf(data):
    return tc.text_to_pdf(pptx_to_text(data))

def pptx_to_images(data, fmt="png"):
    prs = Presentation(io.BytesIO(data))
    # python-pptx does not render slides itself. Return a clear error rather than fake conversion.
    raise ValueError("PPTX image rendering requires a slide rendering engine (such as LibreOffice).")
