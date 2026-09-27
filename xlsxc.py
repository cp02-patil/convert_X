import io, csv, html
from openpyxl import load_workbook
import xlsxc as tc

def _rows(data):
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    for ws in wb.worksheets:
        yield ws.title, list(ws.iter_rows(values_only=True))
    wb.close()

def xlsx_to_csv(data):
    sheets = list(_rows(data))
    out = io.StringIO()
    w = csv.writer(out)
    for idx, (name, rows) in enumerate(sheets):
        if idx: w.writerow([])
        w.writerow([f"[Sheet: {name}]"])
        w.writerows(rows)
    return out.getvalue()

def xlsx_to_text(data):
    parts = []
    for name, rows in _rows(data):
        parts.append(f"[Sheet: {name}]")
        parts.extend("\t".join("" if v is None else str(v) for v in r) for r in rows)
    return "\n".join(parts)

def xlsx_to_html(data):
    tables = []
    for name, rows in _rows(data):
        trs = [f"<tr>{''.join('<td>'+html.escape('' if v is None else str(v))+'</td>' for v in r)}</tr>" for r in rows]
        tables.append(f"<h2>{html.escape(name)}</h2><table border='1'>{''.join(trs)}</table>")
    return "<html><body>" + "".join(tables) + "</body></html>"

def xlsx_to_pdf(data):
    return tc.text_to_pdf(xlsx_to_text(data))
